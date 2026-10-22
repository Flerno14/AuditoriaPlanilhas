"""Ponte JSON Lines entre Electron e a lógica de planilhas em Python."""
import base64
import hashlib
from datetime import date, datetime
from io import BytesIO
import json
import sys
import unicodedata
from pathlib import Path
from openpyxl import load_workbook

import pandas as pd
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import app_corrigido as core

REPORT_COLUMNS = [
    "Índice/Linha",
    "Nome da Coluna",
    "Valor no Arquivo 1",
    "Valor no Arquivo 2",
    "Decisão",
]

# One normalized pair only; no workbook bytes, decisions or generated files.
COMPARISON_CACHE_MAX_BYTES = 64 * 1024 * 1024
_comparison_cache = {}


def comparison_frames(first, second, data, fingerprint):
    key = (str(Path(data['file1']).resolve()), str(Path(data['file2']).resolve()), fingerprint)
    frames = _comparison_cache.get(key)
    if frames is not None:
        return frames
    _comparison_cache.clear()
    frames = (
        core.normalizar_dataframe(core.carregar_excel(first, data['sheet'], 'Arquivo 1')),
        core.normalizar_dataframe(core.carregar_excel(second, data['sheet'], 'Arquivo 2')),
    )
    core.validar_colunas(*frames, bool(data.get('alignByFirstColumn', False)))
    size = sum(int(frame.memory_usage(index=True, deep=True).sum())
               + int(frame.columns.memory_usage(deep=True)) for frame in frames)
    if size <= COMPARISON_CACHE_MAX_BYTES:
        _comparison_cache[key] = frames
    return frames


def header_id(value):
    """Compare localized report headings regardless of accent or encoding artifacts."""
    decomposed = unicodedata.normalize("NFKD", str(value).casefold())
    return "".join(char for char in decomposed if char.isalnum() and not unicodedata.combining(char))


REPORT_HEADER_IDS = {header_id(column): column for column in REPORT_COLUMNS}


def report_frame(records):
    """Restore a report with the exact column names expected by the core logic."""
    normalized = []
    for record in records:
        row = {}
        for key, value in restore_value(record).items():
            row[REPORT_HEADER_IDS.get(header_id(key), key)] = value
        normalized.append(row)
    extra_columns = ["Valor personalizado"] if any("Valor personalizado" in row for row in normalized) else []
    frame = pd.DataFrame(normalized, columns=REPORT_COLUMNS + extra_columns)
    for row in normalized:
        if row.get("Decisão") == "Valor personalizado" and not isinstance(row.get("Valor personalizado"), str):
            raise ValueError("A decisão personalizada exige um valor de texto, que pode estar vazio.")
    if normalized:
        missing = [
            column
            for column in REPORT_COLUMNS
            if any(column not in row for row in normalized)
        ]
        if missing:
            raise ValueError(
                "Relatório recebido sem as colunas obrigatórias: " + ", ".join(missing)
            )
    return frame

def json_value(value):
    if value is None or value is pd.NA:
        return None
    ausente = pd.isna(value)
    # pd.NA has no truth value. Values emitted by a DataFrame are scalar, but
    # retaining this guard also avoids treating a collection as a boolean.
    if not hasattr(ausente, "__len__") and bool(ausente):
        return None
    if hasattr(value, "isoformat"):
        return {"__excel_datetime__": value.isoformat(), "date_only": isinstance(value, date) and not isinstance(value, datetime)}
    if hasattr(value, "item"):
        return json_value(value.item())
    if isinstance(value, str):
        return core.sanitizar_texto_excel(value)
    return value


def restore_value(value):
    if isinstance(value, dict) and "__excel_datetime__" in value:
        if value.get("date_only"):
            return date.fromisoformat(value["__excel_datetime__"])
        return datetime.fromisoformat(value["__excel_datetime__"])
    if isinstance(value, dict):
        return {key: restore_value(item) for key, item in value.items()}
    if isinstance(value, list):
        return [restore_value(item) for item in value]
    if isinstance(value, str):
        return core.sanitizar_texto_excel(value)
    return value


def dataframe_rows(frame):
    return [
        {str(key): json_value(value) for key, value in row.items()}
        for row in frame.to_dict(orient="records")
    ]


def export_frame(frame, template_path=None, text_cells=()):
    if not template_path:
        output = BytesIO()
        frame.to_excel(output, index=False, sheet_name="Relatório")
        if not text_cells:
            return output.getvalue()
        workbook = load_workbook(BytesIO(output.getvalue()))
        for row, column in text_cells:
            cell = workbook.worksheets[0].cell(row=row, column=column)
            cell.value = frame.iloc[row - 2, column - 1]
            cell.data_type = "s"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()
    template = Path(template_path)
    if not template.is_file():
        raise ValueError(f"O modelo de exportação não foi encontrado: {template}")
    workbook = load_workbook(template)
    sheet = workbook.worksheets[0]
    for row_index, values in enumerate([list(frame.columns), *frame.itertuples(index=False, name=None)], start=1):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.value = restore_value(json_value(value))
            if (row_index, column_index) in text_cells:
                cell.data_type = "s"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def session_fingerprint(first, second, data):
    aligned = bool(data.get('alignByFirstColumn', False))
    context = [data['sheet'], aligned,
               bool(data.get('validateUniqueFirstColumn', True)) if aligned else False]
    identity = [hashlib.sha256(first).hexdigest(), hashlib.sha256(second).hexdigest(), context]
    return hashlib.sha256(json.dumps(identity, ensure_ascii=True).encode('utf-8')).hexdigest()


def checked_session_files(data):
    try:
        first = Path(data['file1']).read_bytes()
        second = Path(data['file2']).read_bytes()
    except OSError as error:
        _comparison_cache.clear()
        raise ValueError('Não foi possível acessar os arquivos da revisão. Verifique se foram movidos, removidos ou estão indisponíveis.') from error
    if data.get('fingerprint') != session_fingerprint(first, second, data):
        _comparison_cache.clear()
        raise ValueError('Os arquivos ou as regras mudaram desde a comparação. As decisões antigas não podem ser aplicadas. Compare novamente para iniciar uma nova revisão.')
    return first, second


def report_progress(data, message):
    if data.get('reportProgress'):
        sys.stdout.write(json.dumps({'id': data['id'], 'progress': message}, ensure_ascii=True) + '\n')
        sys.stdout.flush()


def execute(action, data):
    report_progress(data, 'Lendo arquivos e verificando a solicitação...')
    if action.startswith('translation-'):
        from python import translation_audit
        return translation_audit.execute(action, data)
    if action == 'row-details':
        first, second = checked_session_files(data)
        frames = comparison_frames(first, second, data, data['fingerprint'])
        details = core.obter_detalhes_linha_frames(
            *frames, data['line'],
            bool(data.get('alignByFirstColumn', False)),
            bool(data.get('validateUniqueFirstColumn', True)),
        )
        for key in ('values1', 'values2'):
            details[key] = [json_value(value) for value in details[key]]
        return details
    if action == 'preview-new-rows':
        first, second = checked_session_files(data)
        report_progress(data, 'Preparando as linhas novas para visualização...')
        df1, df2 = comparison_frames(first, second, data, data['fingerprint'])
        aligned = bool(data.get('alignByFirstColumn', False))
        unique = bool(data.get('validateUniqueFirstColumn', True))
        core.validar_colunas(df1, df2, aligned)
        indices = core.obter_indices_linhas_novas(df1, df2, aligned, unique)
        rows = df2.reindex(columns=df1.columns).iloc[indices]
        return {'columns': list(df1.columns), 'rows': [
            {'line': index + 2, 'values': [json_value(value) for value in values]}
            for index, values in zip(indices, rows.itertuples(index=False, name=None))
        ]}
    if action == 'validate-session':
        checked_session_files(data)
        return True
    if action == "sheets":
        return core.obter_abas(Path(data["path"]).read_bytes())

    if action == "compare":
        _comparison_cache.clear()
        first = Path(data["file1"]).read_bytes()
        second = Path(data["file2"]).read_bytes()
        fingerprint = session_fingerprint(first, second, data)
        report_progress(data, 'Validando a estrutura e comparando os dados da aba...')
        try:
            frames = comparison_frames(first, second, data, fingerprint)
            report, new_rows = core.comparar_frames(
                *frames,
                alinhar_primeira_coluna=bool(data.get("alignByFirstColumn", False)),
                validar_chaves_unicas=bool(data.get("validateUniqueFirstColumn", True)),
            )
        except Exception:
            _comparison_cache.clear()
            raise
        report_progress(data, 'Preparando as diferenças para a revisão...')
        return {"report": dataframe_rows(report), "newRows": new_rows, "fingerprint": fingerprint}

    if action == "export-differences":
        report_progress(data, 'Montando o relatório de diferenças...')
        frame = report_frame(data["report"])
        text_cells = {(index + 2, 6) for index, row in frame.iterrows() if row["Decisão"] == "Valor personalizado"}
        content = export_frame(frame, data.get("templatePath"), text_cells)
        return {"content": base64.b64encode(content).decode("ascii")}

    if action == "export-choices":
        report_progress(data, 'Montando o relatório de escolhas...')
        frame = report_frame(data["report"])
        if frame["Decisão"].eq("Pendente").any():
            raise ValueError("Resolva todas as diferenças antes de exportar as escolhas.")
        selected = frame.apply(lambda row: row["Valor personalizado"] if row["Decisão"] == "Valor personalizado" else row["Valor no Arquivo 2"] if row["Decisão"] == "Usar Arquivo 2" else row["Valor no Arquivo 1"], axis=1)
        text_cells = {(index + 2, 3) for index, row in frame.iterrows() if row["Decisão"] == "Valor personalizado"}
        content = export_frame(pd.DataFrame({"Linha": frame["Índice/Linha"], "Coluna": frame["Nome da Coluna"], "Mudanças Aplicadas": selected}), data.get("templatePath"), text_cells)
        return {"content": base64.b64encode(content).decode("ascii")}

    if action == "generate":
        first, second = checked_session_files(data)
        report_progress(data, 'Aplicando decisões e preparando a cópia do Excel...')
        report = report_frame(data["report"])
        decisions = report.copy()
        content, changes, rows = core.gerar_arquivo_corrigido(
            first, second, report, decisions, data["sheet"], data["name"],
            incluir_linhas_novas=data["includeNewRows"],
            marcar_alteracoes_laranja=data["markChanges"],
            cor_alteracoes=data.get("changeColor", "#FFA500"),
            alinhar_primeira_coluna=bool(data.get("alignByFirstColumn", False)),
            validar_chaves_unicas=bool(data.get("validateUniqueFirstColumn", True)),
        )
        report_progress(data, 'Preparando o arquivo para salvar...')
        return {"content": base64.b64encode(content).decode("ascii"), "changes": changes, "rows": rows}

    raise ValueError(f"Operação desconhecida: {action}")


def main():
    # Electron writes JSON Lines as UTF-8, including localized report keys.
    sys.stdin.reconfigure(encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    for line in sys.stdin:
        try:
            request = {}
            request = json.loads(line)
            result = execute(request["action"], request)
            response = {"id": request["id"], "ok": True, "result": result}
        except Exception as error:
            response = {"id": request.get("id"), "ok": False, "error": str(error)}
        # Keep isolated surrogates safe for the UTF-8 JSON Lines transport.
        sys.stdout.write(json.dumps(response, ensure_ascii=True, default=json_value) + "\n")
        sys.stdout.flush()


if __name__ == '__main__':
    main()
