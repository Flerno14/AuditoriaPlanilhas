import base64
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

import app_corrigido as core
from test_structure import spreadsheet


def generate(first, second, report, **options):
    return core.gerar_arquivo_corrigido(first, second, report, report, 'Dados',
                                       options.pop('name', 'original.xlsx'), **options)


class BusinessRulesTests(unittest.TestCase):
    def test_exports_write_first_template_sheet_even_when_another_is_active(self):
        with tempfile.TemporaryDirectory() as folder:
            template = Path(folder) / 'modelo.xlsx'
            book = load_workbook(BytesIO(spreadsheet(['Modelo'])))
            book.active['A1'].fill = PatternFill('solid', fgColor='FF123456')
            book.create_sheet('Outra')['A1'] = 'Não alterar'
            book.active = 1
            book.save(template)
            book.close()
            original = template.read_bytes()
            report = [{'Índice/Linha': 2, 'Nome da Coluna': 'Data',
                       'Valor no Arquivo 1': None,
                       'Valor no Arquivo 2': {'__excel_datetime__': '2026-10-11T14:30:45', 'date_only': False},
                       'Decisão': 'Usar Arquivo 2'}]
            for action, header, date_cell in [('export-differences', 'Índice/Linha', 'D2'),
                                              ('export-choices', 'Linha', 'C2')]:
                with self.subTest(action=action):
                    request = dict(id=1, action=action, report=report, templatePath=str(template))
                    result = subprocess.run([sys.executable, str(Path(__file__).resolve().parents[1] / 'python/bridge.py')],
                        input=json.dumps(request) + '\n', text=True, encoding='utf-8', capture_output=True, timeout=30)
                    self.assertEqual(result.returncode, 0, result.stderr)
                    response = json.loads(result.stdout)
                    self.assertTrue(response['ok'], response)
                    output = load_workbook(BytesIO(base64.b64decode(response['result']['content'])))
                    self.assertEqual(output.worksheets[0]['A1'].value, header)
                    self.assertEqual(output.worksheets[0]['A2'].value, 2)
                    self.assertEqual(output.worksheets[0][date_cell].value, datetime(2026, 10, 11, 14, 30, 45))
                    self.assertEqual(output.worksheets[0]['A1'].fill.fgColor.rgb, 'FF123456')
                    self.assertEqual(output['Outra']['A1'].value, 'Não alterar')
                    output.close()
                    self.assertEqual(template.read_bytes(), original)

    def test_empty_second_values_keep_original_and_empty_first_can_be_filled(self):
        first = spreadsheet(['ID', 'Valor'], [[1, 'Manter'], [2, None], [3, None], [4, 'Texto']])
        second = spreadsheet(['ID', 'Valor'], [[1, None], [2, 'Novo'], [3, None], [4, '  ']])
        report, count = core.comparar_dataframes(first, second, 'Dados')
        self.assertEqual(report['Índice/Linha'].tolist(), [3])
        self.assertEqual(report['Valor no Arquivo 2'].tolist(), ['Novo'])
        self.assertEqual(count, 0)
        report['Decisão'] = 'Usar Arquivo 2'
        content, changes, _ = generate(first, second, report)
        book = load_workbook(BytesIO(content))
        self.assertEqual([book.active.cell(i, 2).value for i in range(2, 6)], ['Manter', 'Novo', None, 'Texto'])
        self.assertEqual(changes, 1)
        book.close()

    def test_original_only_rows_remain_unchanged(self):
        first = spreadsheet(['ID', 'Valor'], [[1, 'A'], [2, 'B']])
        second = spreadsheet(['ID', 'Valor'], [[1, 'A']])
        report, count = core.comparar_dataframes(first, second, 'Dados')
        self.assertTrue(report.empty)
        self.assertEqual(count, 0)
        content, changes, rows = generate(first, second, report)
        book = load_workbook(BytesIO(content))
        self.assertEqual(list(book.active.values), [('ID', 'Valor'), (1, 'A'), (2, 'B')])
        self.assertEqual((changes, rows), (0, 0))
        book.close()

    def test_new_rows_are_optional_and_not_cell_differences(self):
        first = spreadsheet(['ID', 'Valor'], [[1, 'A']])
        second = spreadsheet(['ID', 'Valor'], [[1, 'A'], [2, 'B']])
        report, count = core.comparar_dataframes(first, second, 'Dados')
        self.assertTrue(report.empty)
        self.assertEqual(count, 1)
        for include in (False, True):
            with self.subTest(include=include):
                content, changes, rows = generate(first, second, report, incluir_linhas_novas=include)
                book = load_workbook(BytesIO(content))
                self.assertEqual(book.active.max_row, 3 if include else 2)
                self.assertEqual((changes, rows), (0, int(include)))
                book.close()

    def test_alignment_tracks_original_row_positions_and_new_keys(self):
        first = spreadsheet(['ID', 'Valor'], [[1, 'A'], [2, 'B'], [3, 'C']])
        second = spreadsheet(['ID', 'Valor'], [[2, 'Atualizado'], [1, 'A'], [4, 'D']])
        report, count = core.comparar_dataframes(first, second, 'Dados', alinhar_primeira_coluna=True)
        self.assertEqual(report['Índice/Linha'].tolist(), [3])
        self.assertEqual(count, 1)
        report['Decisão'] = 'Usar Arquivo 2'
        content, changes, rows = generate(first, second, report, alinhar_primeira_coluna=True)
        book = load_workbook(BytesIO(content))
        self.assertEqual(list(book.active.values)[1:], [(1, 'A'), (2, 'Atualizado'), (3, 'C'), (4, 'D')])
        self.assertEqual((changes, rows), (1, 1))
        book.close()

    def test_repeated_keys_match_in_occurrence_order_when_allowed(self):
        first = spreadsheet(['ID', 'Valor'], [['x', 'A'], ['x', 'B']])
        second = spreadsheet(['ID', 'Valor'], [['x', 'A'], ['x', 'Novo'], ['x', 'Extra']])
        with self.assertRaisesRegex(ValueError, 'chave duplicada'):
            core.comparar_dataframes(first, second, 'Dados', alinhar_primeira_coluna=True)
        report, count = core.comparar_dataframes(first, second, 'Dados', True, False)
        self.assertEqual(report['Índice/Linha'].tolist(), [3])
        self.assertEqual(count, 1)
        report['Decisão'] = 'Usar Arquivo 2'
        content, _, rows = generate(first, second, report, alinhar_primeira_coluna=True, validar_chaves_unicas=False)
        book = load_workbook(BytesIO(content))
        self.assertEqual([row[1] for row in list(book.active.values)[1:]], ['A', 'Novo', 'Extra'])
        self.assertEqual(rows, 1)
        book.close()

    def test_blank_keys_rejected_even_when_duplicates_allowed(self):
        valid = spreadsheet(['ID', 'Valor'], [[1, 'A']])
        for value in (None, '  '):
            invalid = spreadsheet(['ID', 'Valor'], [[value, 'A']])
            for first, second in ((valid, invalid), (invalid, valid)):
                with self.subTest(value=value, first_valid=first == valid):
                    with self.assertRaisesRegex(ValueError, 'chave vazia'):
                        core.comparar_dataframes(first, second, 'Dados', True, False)

    def test_only_approved_changes_are_written_and_highlighted(self):
        first = spreadsheet(['ID', 'Valor'], [[1, 'A'], [2, 'B']])
        second = spreadsheet(['ID', 'Valor'], [[1, 'Novo A'], [2, 'Novo B']])
        report, _ = core.comparar_dataframes(first, second, 'Dados')
        report['Decisão'] = ['Usar Arquivo 2', 'Manter Arquivo 1']
        for mark in (False, True):
            content, changes, _ = generate(first, second, report, marcar_alteracoes_laranja=mark, cor_alteracoes='#123456')
            book = load_workbook(BytesIO(content))
            self.assertEqual(book.active['B2'].value, 'Novo A')
            self.assertEqual(book.active['B3'].value, 'B')
            self.assertEqual(changes, 1)
            self.assertIsNone(book.active['B3'].fill.patternType)
            if mark:
                self.assertEqual(book.active['B2'].fill.fgColor.rgb, 'FF123456')
            else:
                self.assertIsNone(book.active['B2'].fill.patternType)
            book.close()

    def test_untouched_formulas_styles_and_other_sheets_are_preserved(self):
        raw = spreadsheet(['ID', 'Valor', 'Total'], [[1, 'A', '=1+2']])
        book = load_workbook(BytesIO(raw))
        book.active['B2'].fill = PatternFill('solid', fgColor='FF112233')
        book.active['C2'].number_format = '0.00'
        book.create_sheet('Auxiliar')['A1'] = '=SUM(Dados!A2)'
        output = BytesIO(); book.save(output); book.close()
        first = output.getvalue()
        second = spreadsheet(['ID', 'Valor', 'Total'], [[1, 'B', '=1+2']])
        report, _ = core.comparar_dataframes(first, second, 'Dados')
        report['Decisão'] = 'Usar Arquivo 2'
        content, _, _ = generate(first, second, report)
        book = load_workbook(BytesIO(content), data_only=False)
        self.assertEqual(book.active['C2'].value, '=1+2')
        self.assertEqual(book.active['C2'].number_format, '0.00')
        self.assertEqual(book.active['B2'].fill.fgColor.rgb, 'FF112233')
        self.assertEqual(book['Auxiliar']['A1'].value, '=SUM(Dados!A2)')
        book.close()

    def test_missing_values_and_mixed_types_keep_excel_types(self):
        for value in (None, pd.NA, np.nan, pd.NaT):
            self.assertIsNone(core.normalizar_valor_excel(value))
        first = spreadsheet(['ID', 'Valor'], [[1, 'old'], [2, 'old'], [3, 'old'], [4, 'old']])
        expected = [123, '00123', 'N/A', datetime(2026, 2, 3, 14, 30)]
        second = spreadsheet(['ID', 'Valor'], list(enumerate(expected, start=1)))
        report, _ = core.comparar_dataframes(first, second, 'Dados')
        report['Decisão'] = 'Usar Arquivo 2'
        content, _, _ = generate(first, second, report)
        book = load_workbook(BytesIO(content))
        self.assertEqual([book.active.cell(i, 2).value for i in range(2, 6)], expected)
        book.close()

    def test_xlsm_preserves_vba_part_bytes(self):
        # Synthetic package checks byte preservation only; no executable macro.
        first = spreadsheet(['ID', 'Valor'], [[1, 'A']])
        marker = b'Synthetic VBA preservation fixture'
        output = BytesIO()
        with ZipFile(BytesIO(first)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
            for name in source.namelist():
                content = source.read(name)
                if name == '[Content_Types].xml':
                    content = content.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml', b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
                    content = content.replace(b'</Types>', b'<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
                if name == 'xl/_rels/workbook.xml.rels':
                    content = content.replace(b'</Relationships>', b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
                target.writestr(name, content)
            target.writestr('xl/vbaProject.bin', marker)
        first = output.getvalue()
        second = spreadsheet(['ID', 'Valor'], [[1, 'B']])
        report, _ = core.comparar_dataframes(first, second, 'Dados')
        report['Decisão'] = 'Usar Arquivo 2'
        content, _, _ = generate(first, second, report, name='original.xlsm')
        with ZipFile(BytesIO(content)) as result:
            self.assertEqual(result.read('xl/vbaProject.bin'), marker)
            self.assertIn(b'macroEnabled', result.read('[Content_Types].xml'))

    def test_bridge_roundtrip_dates_and_pending_export_rejection(self):
        with tempfile.TemporaryDirectory() as folder:
            first, second = Path(folder) / '1.xlsx', Path(folder) / '2.xlsx'
            first.write_bytes(spreadsheet(['ID', 'Data'], [[1, datetime(2025, 1, 1)]]))
            second.write_bytes(spreadsheet(['ID', 'Data'], [[1, datetime(2026, 2, 3)]]))
            def request(action, **data):
                result = subprocess.run([sys.executable, str(Path(__file__).resolve().parents[1] / 'python/bridge.py')],
                    input=json.dumps(dict(id=1, action=action, **data)) + '\n', text=True, encoding='utf-8', capture_output=True, timeout=30)
                self.assertEqual(result.returncode, 0, result.stderr)
                return json.loads(result.stdout)
            context = dict(file1=str(first), file2=str(second), sheet='Dados')
            compared = request('compare', **context)['result']
            report = compared['report']
            self.assertIn('__excel_datetime__', report[0]['Valor no Arquivo 2'])
            rejected = request('export-choices', report=report)
            self.assertFalse(rejected['ok'])
            self.assertIn('Resolva todas', rejected['error'])
            report[0]['Decisão'] = 'Usar Arquivo 2'
            generated = request('generate', **context, fingerprint=compared['fingerprint'], report=report,
                                name='original.xlsx', includeNewRows=True, markChanges=False)
            self.assertTrue(generated['ok'], generated)
            book = load_workbook(BytesIO(base64.b64decode(generated['result']['content'])))
            self.assertEqual(book.active['B2'].value, datetime(2026, 2, 3))
            book.close()
