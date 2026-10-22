from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
import pandas as pd

import app_corrigido as core


def spreadsheet(headers, rows=()):
    book = Workbook()
    book.active.title = "Dados"
    if headers:
        book.active.append(headers)
    for row in rows:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class StructureTests(unittest.TestCase):
    def test_extra_columns_block_even_without_values_or_shared_rows(self):
        first = spreadsheet(["ID", "Nome"])
        for rows in ([], [[1, "Ana", None]], [[1, "Ana", "Extra"]]):
            with self.subTest(rows=rows):
                second = spreadsheet(["ID", "Nome", "Observação"], rows)
                with self.assertRaisesRegex(ValueError, "Colunas exclusivas do Arquivo 2: Observação"):
                    core.comparar_dataframes(first, second, "Dados")

    def test_incompatible_names_list_both_sides(self):
        first = spreadsheet(["ID", "Descrição"])
        second = spreadsheet(["ID", "Descricao"])
        with self.assertRaises(ValueError) as error:
            core.comparar_dataframes(first, second, "Dados")
        self.assertIn("Arquivo 2: Descricao", str(error.exception))
        self.assertIn("Arquivo 1: Descrição", str(error.exception))

    def test_duplicate_original_headers_are_not_hidden_by_pandas(self):
        valid = spreadsheet(["ID", "Nome"], [[1, "Ana"]])
        for headers in (["ID", "ID"], ["ID", " ID "], [1, "1"]):
            for slot in (1, 2):
                with self.subTest(headers=headers, slot=slot):
                    invalid = spreadsheet(headers, [[1, 2]])
                    first, second = (invalid, valid) if slot == 1 else (valid, invalid)
                    with self.assertRaises(ValueError) as error:
                        core.comparar_dataframes(first, second, "Dados")
                    self.assertIn(f"Arquivo {slot}", str(error.exception))
                    self.assertIn("Cabeçalhos duplicados", str(error.exception))

    def test_empty_headers_including_trailing_column_with_data(self):
        valid = spreadsheet(["ID", "Nome"], [[1, "Ana"]])
        for headers in (["ID", None], ["ID", "  "], [None, "Nome"]):
            with self.subTest(headers=headers):
                invalid = spreadsheet(headers, [[1, "Ana"]])
                with self.assertRaisesRegex(ValueError, "Cabeçalhos vazios"):
                    core.comparar_dataframes(valid, invalid, "Dados")

    def test_empty_sheet_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "Arquivo 1.*não possui cabeçalhos"):
            core.comparar_dataframes(spreadsheet([]), spreadsheet(["ID"]), "Dados")

    def test_header_only_sheet_is_valid(self):
        report, count = core.comparar_dataframes(spreadsheet(["ID"]), spreadsheet(["ID"]), "Dados")
        self.assertTrue(report.empty)
        self.assertEqual(count, 0)

    def test_reordered_columns_are_written_to_correct_original_positions(self):
        first = spreadsheet(["ID", "Nome"], [[1, "Ana"]])
        second = spreadsheet([" Nome ", "ID"], [["Bia", 1]])
        report, _ = core.comparar_dataframes(first, second, "Dados")
        self.assertEqual(report["Nome da Coluna"].tolist(), ["Nome"])
        report["Decisão"] = "Usar Arquivo 2"
        content, changes, _ = core.gerar_arquivo_corrigido(first, second, report, report, "Dados", "original.xlsx")
        book = load_workbook(BytesIO(content))
        self.assertEqual(list(book.active.values), [("ID", "Nome"), (1, "Bia")])
        self.assertEqual(changes, 1)
        book.close()

    def test_columns_only_in_original_remain_intact(self):
        first = spreadsheet(["ID", "Nome", "Nota"], [[1, "Ana", "Preservar"]])
        second = spreadsheet(["ID", "Nome"], [[1, "Bia"], [2, "Caio"]])
        report, count = core.comparar_dataframes(first, second, "Dados")
        self.assertEqual(count, 1)
        self.assertEqual(report["Nome da Coluna"].tolist(), ["Nome"])
        report["Decisão"] = "Usar Arquivo 2"
        content, _, rows = core.gerar_arquivo_corrigido(first, second, report, report, "Dados", "original.xlsx")
        book = load_workbook(BytesIO(content))
        self.assertEqual(book.active["C2"].value, "Preservar")
        self.assertEqual(book.active["B3"].value, "Caio")
        self.assertIsNone(book.active["C3"].value)
        self.assertEqual(rows, 1)
        book.close()

    def test_alignment_rejects_different_first_column(self):
        first = spreadsheet(["ID", "Nome"], [[1, "Ana"]])
        second = spreadsheet(["Nome", "ID"], [["Ana", 1]])
        with self.assertRaisesRegex(ValueError, "Alinhamento incompatível"):
            core.comparar_dataframes(first, second, "Dados", alinhar_primeira_coluna=True)

    def test_alignment_accepts_reordered_non_key_columns(self):
        first = spreadsheet(["ID", "Nome", "Nota"], [[1, "Ana", "A"], [2, "Bia", "B"]])
        second = spreadsheet([" ID ", "Nota", "Nome"], [[2, "B", "Bia"], [1, "A", "Ana"]])
        report, count = core.comparar_dataframes(first, second, "Dados", alinhar_primeira_coluna=True)
        self.assertTrue(report.empty)
        self.assertEqual(count, 0)

    def test_generation_also_rejects_incompatible_structure(self):
        first = spreadsheet(["ID"], [[1]])
        second = spreadsheet(["ID", "Extra"], [[1, "X"]])
        with self.assertRaisesRegex(ValueError, "Estruturas incompatíveis"):
            core.gerar_arquivo_corrigido(first, second, pd.DataFrame(), pd.DataFrame(), "Dados", "original.xlsx")

    def test_literal_unnamed_header_is_valid(self):
        data = spreadsheet(["Unnamed: 0", "NA"], [["N/A", "NA"]])
        frame = core.carregar_excel(data, "Dados")
        self.assertEqual(frame.iloc[0].tolist(), ["N/A", "NA"])

    def test_separated_notes_are_ignored_and_preserved_in_generated_workbook(self):
        first = spreadsheet(['ID', 'Nome', None, None], [[1, 'Ana', None, 'Anotação']])
        second = spreadsheet(['ID', 'Nome'], [[1, 'Bia']])
        report, count = core.comparar_dataframes(first, second, 'Dados')
        self.assertEqual(report['Nome da Coluna'].tolist(), ['Nome'])
        self.assertEqual(count, 0)
        report['Decisão'] = 'Usar Arquivo 2'
        content, changes, _ = core.gerar_arquivo_corrigido(first, second, report, report, 'Dados', 'original.xlsx')
        book = load_workbook(BytesIO(content))
        self.assertEqual(book.active['B2'].value, 'Bia')
        self.assertEqual(book.active['D2'].value, 'Anotação')
        self.assertEqual(changes, 1)
        book.close()

    def test_separated_notes_in_second_file_do_not_become_new_columns(self):
        first = spreadsheet(['ID', 'Nome'], [[1, 'Ana']])
        second = spreadsheet(['ID', 'Nome', None, None], [[1, 'Ana', None, 'Nota']])
        report, count = core.comparar_dataframes(first, second, 'Dados')
        self.assertTrue(report.empty)
        self.assertEqual(count, 0)

    def test_blank_separator_before_named_column_still_requires_header(self):
        data = spreadsheet(['ID', None, 'Nome'], [[1, None, 'Ana']])
        with self.assertRaisesRegex(ValueError, 'Cabeçalhos vazios'):
            core.carregar_excel(data, 'Dados')

    def test_bridge_returns_validation_error_as_json_line(self):
        with tempfile.TemporaryDirectory() as folder:
            first = Path(folder) / "original.xlsx"
            second = Path(folder) / "modificado.xlsx"
            first.write_bytes(spreadsheet(["ID"]))
            second.write_bytes(spreadsheet(["ID", "Extra"]))
            request = {"id": 1, "action": "compare", "file1": str(first), "file2": str(second), "sheet": "Dados"}
            result = subprocess.run(
                [sys.executable, str(Path(__file__).resolve().parents[1] / "python" / "bridge.py")],
                input=json.dumps(request) + "\n", capture_output=True, text=True, encoding="utf-8", timeout=30,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            response = json.loads(result.stdout)
            self.assertEqual(response["id"], 1)
            self.assertFalse(response["ok"])
            self.assertIn("Colunas exclusivas do Arquivo 2: Extra", response["error"])


if __name__ == "__main__":
    unittest.main()
