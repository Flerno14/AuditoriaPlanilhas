import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from test_structure import spreadsheet


class SessionIdentityTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.first = Path(self.folder.name) / 'original.xlsx'
        self.second = Path(self.folder.name) / 'modificado.xlsx'
        self.first.write_bytes(spreadsheet(['ID', 'Nome'], [[1, 'Ana']]))
        self.second.write_bytes(spreadsheet(['ID', 'Nome'], [[1, 'Bia']]))
        self.context = dict(file1=str(self.first), file2=str(self.second), sheet='Dados',
                            alignByFirstColumn=False, validateUniqueFirstColumn=True)
        self.compared = self.request('compare', self.context)['result']
        self.context['fingerprint'] = self.compared['fingerprint']

    def request(self, action, data):
        result = subprocess.run(
            [sys.executable, str(Path(__file__).resolve().parents[1] / 'python' / 'bridge.py')],
            input=json.dumps(dict(data, id=1, action=action)) + '\n',
            text=True, encoding='utf-8', capture_output=True, timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        return json.loads(result.stdout)

    def generation(self):
        report = self.compared['report']
        for row in report:
            row['Decisão'] = 'Usar Arquivo 2'
        return dict(self.context, report=report, name='original.xlsx', includeNewRows=True, markChanges=False)

    def test_unchanged_files_resume_and_generate(self):
        self.assertTrue(self.request('validate-session', self.context)['ok'])
        self.assertTrue(self.request('generate', self.generation())['ok'])
        inactive_rule = dict(self.context, validateUniqueFirstColumn=False)
        self.assertTrue(self.request('validate-session', inactive_rule)['ok'])

    def test_custom_values_generate_and_export_as_text(self):
        import base64
        from io import BytesIO
        from openpyxl import load_workbook
        data = self.generation()
        row = data['report'][0]
        row['Decisão'] = 'Valor personalizado'
        invalid = self.request('generate', data)
        self.assertFalse(invalid['ok'])
        for value in ['001\nJoão', '=1+2', '']:
            row['Valor personalizado'] = value
            for marked in [False, True]:
                response = self.request('generate', dict(data, markChanges=marked))
                self.assertTrue(response['ok'], response)
                book = load_workbook(BytesIO(base64.b64decode(response['result']['content'])))
                self.assertEqual(book['Dados']['B2'].value, value or None)
                self.assertNotEqual(book['Dados']['B2'].data_type, 'f')
                self.assertEqual(book['Dados']['B2'].fill.patternType, 'solid' if marked else None)
                self.assertEqual(response['result']['changes'], 1)
                book.close()
            for action, column in [('export-choices', 'C'), ('export-differences', 'F')]:
                for template in [None, str(self.first)]:
                    response = self.request(action, dict(report=data['report'], templatePath=template))
                    self.assertTrue(response['ok'], response)
                    book = load_workbook(BytesIO(base64.b64decode(response['result']['content'])))
                    cell = book.worksheets[0][column + '2']
                    self.assertEqual(cell.value, value or None)
                    self.assertNotEqual(cell.data_type, 'f')
                    book.close()

    def test_changed_source_blocks_resume_and_generation(self):
        self.second.write_bytes(spreadsheet(['ID', 'Nome'], [[1, 'Caio']]))
        for action, data in [('validate-session', self.context), ('generate', self.generation())]:
            response = self.request(action, data)
            self.assertFalse(response['ok'])
            self.assertIn('mudaram desde a comparação', response['error'])

    def test_missing_file_blocks_resume(self):
        self.first.unlink()
        response = self.request('validate-session', self.context)
        self.assertFalse(response['ok'])
        self.assertIn('acessar os arquivos', response['error'])

    def test_changed_rules_and_missing_identity_are_rejected(self):
        for context in [dict(self.context, alignByFirstColumn=True), dict(self.context, fingerprint=None)]:
            response = self.request('validate-session', context)
            self.assertFalse(response['ok'])

    def test_preview_matches_generation_for_repeated_keys_and_original_columns(self):
        import base64
        from io import BytesIO
        from openpyxl import load_workbook
        self.first.write_bytes(spreadsheet(['ID', 'Nome', 'Nota'], [[1, 'Ana', 'Original']]))
        self.second.write_bytes(spreadsheet(['ID', 'Nome'], [[2, 'Bia'], [1, 'Ana'], [1, 'Caio']]))
        self.context.update(alignByFirstColumn=True, validateUniqueFirstColumn=False)
        self.compared = self.request('compare', self.context)['result']
        self.context['fingerprint'] = self.compared['fingerprint']
        preview = self.request('preview-new-rows', self.context)['result']
        self.assertEqual(preview['columns'], ['ID', 'Nome', 'Nota'])
        self.assertEqual([row['line'] for row in preview['rows']], [2, 4])
        values = [row['values'] for row in preview['rows']]
        self.assertEqual(values, [[2, 'Bia', None], [1, 'Caio', None]])
        generated = self.request('generate', self.generation())['result']
        book = load_workbook(BytesIO(base64.b64decode(generated['content'])))
        self.assertEqual([list(row) for row in list(book.active.values)[2:]], values)
        book.close()

    def test_preview_by_position_and_modified_file_rejection(self):
        self.second.write_bytes(spreadsheet(['ID', 'Nome'], [[1, 'Bia'], [2, 'Caio']]))
        self.compared = self.request('compare', self.context)['result']
        self.context['fingerprint'] = self.compared['fingerprint']
        preview = self.request('preview-new-rows', self.context)['result']
        self.assertEqual(preview['rows'], [{'line': 3, 'values': [2, 'Caio']}])
        self.second.write_bytes(spreadsheet(['ID', 'Nome'], [[1, 'Mudou']]))
        self.assertFalse(self.request('preview-new-rows', self.context)['ok'])

    def test_opt_in_progress_remains_valid_json_lines_with_final_response(self):
        request = dict(self.context, id=9, action='compare', reportProgress=True)
        result = subprocess.run(
            [sys.executable, str(Path(__file__).resolve().parents[1] / 'python' / 'bridge.py')],
            input=json.dumps(request) + '\n', text=True, encoding='utf-8', capture_output=True, timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        messages = [json.loads(line) for line in result.stdout.splitlines()]
        self.assertGreaterEqual(len(messages), 3)
        self.assertTrue(all(message['id'] == 9 for message in messages))
        self.assertTrue(all('progress' in message for message in messages[:-1]))
        self.assertTrue(messages[-1]['ok'])


if __name__ == '__main__':
    unittest.main()
