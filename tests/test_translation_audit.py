import base64
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import numpy as np
from openpyxl import Workbook, load_workbook
from python import bridge
from python import translation_audit


class TranslationAuditTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.path = Path(self.folder.name) / 'sample.xlsx'
        book = Workbook()
        book.active.append(['ID', 'English', 'Português'])
        book.active.append(['001', 'Hello {name}', 'Olá {name}'])
        book.active.append(['002', 'Missing translation', None])
        book.active.append(['003', 'NA', 'N/A'])
        book.save(self.path)
        book.close()
        self.data = {'path': str(self.path), 'sheet': 'Sheet'}

    def test_reads_all_pairs_and_verifies_identity(self):
        info = bridge.execute('translation-columns', self.data)
        data = {**self.data, 'source': 'English', 'target': 'Português', 'identifier': 'ID', 'fingerprint': info['fingerprint']}
        result = bridge.execute('translation-read', data)
        self.assertEqual(len(result['rows']), 3)
        self.assertEqual(result['rows'][1]['target'], '')
        self.assertEqual(result['rows'][2]['source'], 'NA')
        self.assertEqual(result['rows'][0]['identifier'], '001')
        self.assertEqual(result['rows'][1]['line'], 3)
        with self.assertRaises(ValueError):
            bridge.execute('translation-read', {**data, 'target': 'English'})
        with self.assertRaises(ValueError):
            bridge.execute('translation-read', {**data, 'fingerprint': '0' * 64})

    def test_export_keeps_formula_like_text_literal(self):
        row = {'line': 2, 'identifier': '=1+1', 'source': '=SUM(A1)', 'target': '=teste',
               'score': None, 'priority': 'Alta', 'issues': ['Marcador ausente'], 'review': 'Não revisado'}
        result = bridge.execute('translation-export', {'rows': [row]})
        book = load_workbook(BytesIO(base64.b64decode(result['content'])))
        for address in ['B2', 'C2', 'D2']:
            self.assertEqual(book.active[address].data_type, 's')
        self.assertIsNone(book.active['E2'].value)
        book.close()

    def test_semantic_requires_local_model_without_downloading(self):
        with self.assertRaisesRegex(ValueError, 'pasta local'):
            bridge.execute('translation-semantic', {'modelPath': self.folder.name, 'rows': []})

    def test_generate_changes_only_selected_translation_cells(self):
        book = load_workbook(self.path)
        book.active['B3'] = '=1+1'
        book.active['C2'].number_format = '@'
        book.create_sheet('Other')['A1'] = 'Preserved'
        book.save(self.path)
        book.close()
        original = self.path.read_bytes()
        info = bridge.execute('translation-columns', self.data)
        data = {**self.data, 'source': 'English', 'target': 'Português', 'fingerprint': info['fingerprint'],
                'edits': [{'line': 2, 'replacement': '=literal'}, {'line': 4, 'replacement': ''}]}
        result = bridge.execute('translation-generate', data)
        book = load_workbook(BytesIO(base64.b64decode(result['content'])))
        self.assertEqual(book['Sheet']['C2'].value, '=literal')
        self.assertEqual(book['Sheet']['C2'].data_type, 's')
        self.assertEqual(book['Sheet']['C2'].number_format, '@')
        self.assertIsNone(book['Sheet']['C4'].value)
        self.assertEqual(book['Sheet']['B3'].value, '=1+1')
        self.assertEqual(book['Other']['A1'].value, 'Preserved')
        self.assertEqual(self.path.read_bytes(), original)
        book.close()
        for override in [{'fingerprint': '0' * 64}, {'target': 'English'},
                         {'edits': [{'line': 99, 'replacement': 'bad'}]},
                         {'edits': [{'line': 2, 'replacement': None}]}]:
            with self.assertRaises(ValueError):
                bridge.execute('translation-generate', {**data, **override})

    def test_export_includes_explicit_empty_replacement(self):
        row = {'line': 2, 'identifier': '', 'source': 'Hello', 'target': 'Olá', 'replacement': '',
               'score': None, 'priority': 'Alta', 'issues': [], 'review': 'Não revisado'}
        result = bridge.execute('translation-export', {'rows': [row]})
        book = load_workbook(BytesIO(base64.b64decode(result['content'])))
        self.assertEqual(book.active['D2'].value, 'Olá')
        self.assertEqual(book.active['J2'].value, 'Sim')
        self.assertIsNone(book.active['K2'].value)
        book.close()

    def test_semantic_skips_empty_and_over_limit_without_truncating(self):
        folder = Path(self.folder.name)
        (folder / 'modules.json').write_text('[]')
        class Tokenizer:
            def encode(self, text, truncation):
                self.assert_no_truncation = truncation is False
                return text.split()
        class Model:
            max_seq_length = 4
            tokenizer = Tokenizer()
            def encode(self, texts, **kwargs):
                self.texts = texts
                return np.array([[1., 0.], [0.8, 0.6]])
        model = Model()
        with patch.object(translation_audit, '_model', model), patch.object(translation_audit, '_model_path', str(folder.resolve())):
            result = translation_audit.semantic({'modelPath': str(folder), 'rows': [
                {'source': 'Hello', 'target': 'Olá'}, {'source': 'Hello', 'target': ''},
                {'source': 'one two three four five', 'target': 'Texto'}]})
        self.assertAlmostEqual(result[0]['score'], 0.8)
        self.assertIsNone(result[1]['score'])
        self.assertIsNone(result[2]['score'])
        self.assertIn('limite', result[2]['semanticIssue'])
        self.assertEqual(model.texts, ['Hello', 'Olá'])
        self.assertTrue(model.tokenizer.assert_no_truncation)
