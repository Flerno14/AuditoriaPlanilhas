from io import BytesIO
from pathlib import Path
import hashlib

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import load_workbook


# ============================================================
# CONFIGURAÇÃO DA APLICAÇÃO
# ============================================================

st.set_page_config(
    page_title="Auditoria de Planilhas",
    page_icon="📊",
    layout="wide",
)


# ============================================================
# ESTILO VISUAL
# ============================================================

st.markdown(
    """
    <style>

    .main-title {
        font-size: 32px;
        font-weight: 700;
        margin-bottom: 5px;
    }

    .subtitle {
        color: #666666;
        font-size: 16px;
        margin-bottom: 25px;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# FUNÇÕES DE LEITURA
# ============================================================

@st.cache_data(show_spinner=False)
def obter_abas(arquivo_bytes: bytes) -> list[str]:
    """Obtém os nomes das abas existentes em um arquivo Excel."""

    workbook = load_workbook(
        BytesIO(arquivo_bytes),
        read_only=True,
        data_only=False,
    )

    try:
        return workbook.sheetnames
    finally:
        workbook.close()


@st.cache_data(show_spinner=False)
def carregar_excel(
    arquivo_bytes: bytes,
    sheet_name: str,
) -> pd.DataFrame:
    """Lê uma aba do Excel utilizando pandas."""

    return pd.read_excel(
        BytesIO(arquivo_bytes),
        sheet_name=sheet_name,
        engine="openpyxl",
    )


# ============================================================
# NORMALIZAÇÃO
# ============================================================

def normalizar_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """
    Normaliza o DataFrame antes da comparação.

    - Reseta o índice.
    - Remove espaços dos nomes das colunas.
    - Converte nomes de colunas para string.
    - Trata strings vazias como células vazias.
    """

    resultado = df.copy()

    resultado.reset_index(drop=True, inplace=True)

    resultado.columns = [
        str(coluna).strip()
        for coluna in resultado.columns
    ]

    resultado.replace(
        r"^\s*$",
        pd.NA,
        regex=True,
        inplace=True,
    )

    return resultado


# ============================================================
# VALIDAÇÃO DAS COLUNAS
# ============================================================

def validar_colunas(
    df1: pd.DataFrame,
    df2: pd.DataFrame,
) -> None:
    """Garante que os nomes das colunas não sejam duplicados."""

    duplicadas_1 = df1.columns[df1.columns.duplicated()].tolist()
    duplicadas_2 = df2.columns[df2.columns.duplicated()].tolist()

    if duplicadas_1:
        raise ValueError(
            "O Arquivo 1 possui colunas duplicadas: "
            + ", ".join(map(str, duplicadas_1))
        )

    if duplicadas_2:
        raise ValueError(
            "O Arquivo 2 possui colunas duplicadas: "
            + ", ".join(map(str, duplicadas_2))
        )


# ============================================================
# COMPARAÇÃO
# ============================================================

@st.cache_data(show_spinner=False)
def comparar_dataframes(
    arquivo_1_bytes: bytes,
    arquivo_2_bytes: bytes,
    sheet_name: str,
) -> tuple[pd.DataFrame, int]:
    """
    Compara os dois arquivos de forma vetorizada.

    Regras:
    - Diferenças entre linhas que existem no Arquivo 1 aparecem na interface.
    - Linhas que existem somente no Arquivo 2 NÃO aparecem na interface.
    - As linhas exclusivas do Arquivo 2 são tratadas automaticamente como
      "Usar Arquivo 2" e serão adicionadas ao arquivo corrigido.

    Retorna:
        relatório de diferenças entre as linhas existentes no Arquivo 1
        quantidade de linhas novas existentes somente no Arquivo 2
    """

    df1 = normalizar_dataframe(
        carregar_excel(arquivo_1_bytes, sheet_name)
    )

    df2 = normalizar_dataframe(
        carregar_excel(arquivo_2_bytes, sheet_name)
    )

    validar_colunas(df1, df2)

    quantidade_linhas_arquivo_1 = len(df1)
    quantidade_linhas_arquivo_2 = len(df2)

    quantidade_linhas_novas = max(
        0,
        quantidade_linhas_arquivo_2 - quantidade_linhas_arquivo_1,
    )

    # Mantemos a comparação do tamanho máximo para continuar detectando
    # também linhas que existam somente no Arquivo 1 (possíveis remoções).
    quantidade_linhas = max(
        quantidade_linhas_arquivo_1,
        quantidade_linhas_arquivo_2,
    )

    indices = pd.RangeIndex(
        start=0,
        stop=quantidade_linhas,
    )

    todas_colunas = df1.columns.union(
        df2.columns,
        sort=False,
    )

    alinhado_1 = df1.reindex(
        index=indices,
        columns=todas_colunas,
    )

    alinhado_2 = df2.reindex(
        index=indices,
        columns=todas_colunas,
    )

    iguais = alinhado_1.eq(alinhado_2)

    ambos_vazios = (
        alinhado_1.isna()
        & alinhado_2.isna()
    )

    diferenca = ~(iguais | ambos_vazios)
    diferenca = diferenca.fillna(True)

    linhas, colunas = np.where(diferenca.to_numpy())

    # IMPORTANTÍSSIMO:
    # uma linha que existe somente no Arquivo 2 é uma inclusão automática,
    # portanto não deve aparecer como diferença célula a célula na interface.
    # Já uma linha que existe somente no Arquivo 1 continua sendo exibida,
    # pois representa uma possível remoção.
    somente_linhas_do_arquivo_1 = linhas < quantidade_linhas_arquivo_1

    linhas = linhas[somente_linhas_do_arquivo_1]
    colunas = colunas[somente_linhas_do_arquivo_1]

    if len(linhas) == 0:
        return (
            pd.DataFrame(
                columns=[
                    "Índice/Linha",
                    "Nome da Coluna",
                    "Valor no Arquivo 1",
                    "Valor no Arquivo 2",
                    "Decisão",
                ]
            ),
            quantidade_linhas_novas,
        )

    relatorio = pd.DataFrame(
        {
            "Índice/Linha": linhas + 2,
            "Nome da Coluna": todas_colunas.to_numpy()[colunas],
            "Valor no Arquivo 1": alinhado_1.to_numpy()[
                linhas, colunas
            ],
            "Valor no Arquivo 2": alinhado_2.to_numpy()[
                linhas, colunas
            ],
            "Decisão": "Pendente",
        }
    )

    return relatorio, quantidade_linhas_novas


# ============================================================
# CONVERSÃO DE VALORES
# ============================================================

def normalizar_valor_excel(valor):
    """Converte valores pandas para valores aceitos pelo openpyxl."""

    if valor is None:
        return None

    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()

    if isinstance(valor, np.generic):
        return valor.item()

    return valor


# ============================================================
# GERAÇÃO DO ARQUIVO CORRIGIDO
# ============================================================

def copiar_estilo_linha(worksheet, linha_origem: int, linha_destino: int) -> None:
    """Copia a formatação da última linha de dados para uma nova linha."""

    from copy import copy

    if linha_origem < 1 or linha_origem > worksheet.max_row:
        return

    for coluna in range(1, worksheet.max_column + 1):
        origem = worksheet.cell(
            row=linha_origem,
            column=coluna,
        )
        destino = worksheet.cell(
            row=linha_destino,
            column=coluna,
        )

        if origem.has_style:
            destino._style = copy(origem._style)
        if origem.number_format:
            destino.number_format = origem.number_format
        if origem.alignment:
            destino.alignment = copy(origem.alignment)
        if origem.protection:
            destino.protection = copy(origem.protection)


def gerar_arquivo_corrigido(
    arquivo_1_bytes: bytes,
    arquivo_2_bytes: bytes,
    relatorio_original: pd.DataFrame,
    decisoes: pd.DataFrame,
    sheet_name: str,
    nome_arquivo_original: str,
) -> tuple[bytes, int, int]:
    """
    Cria uma cópia do Arquivo 1, aplica as decisões do usuário e
    adiciona automaticamente ao final as linhas que existem somente
    no Arquivo 2.

    Retorna:
        bytes do arquivo corrigido
        quantidade de células alteradas por decisão do usuário
        quantidade de linhas novas adicionadas automaticamente
    """

    extensao = Path(nome_arquivo_original).suffix.lower()

    parametros = {"data_only": False}

    if extensao == ".xlsm":
        parametros["keep_vba"] = True

    workbook = load_workbook(
        BytesIO(arquivo_1_bytes),
        **parametros,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"A aba '{sheet_name}' não existe no Arquivo 1."
            )

        worksheet = workbook[sheet_name]

        # DataFrames normalizados usados para descobrir o tamanho real
        # das áreas de dados e os valores das linhas novas.
        df_original = normalizar_dataframe(
            carregar_excel(arquivo_1_bytes, sheet_name)
        )

        df_modificado = normalizar_dataframe(
            carregar_excel(arquivo_2_bytes, sheet_name)
        )

        mapa_colunas = {
            str(nome): indice + 1
            for indice, nome in enumerate(df_original.columns)
        }

        # Valores do Arquivo 2 para cada diferença encontrada.
        valores_arquivo_2 = {}

        for _, linha in relatorio_original.iterrows():
            chave = (
                int(linha["Índice/Linha"]),
                str(linha["Nome da Coluna"]),
            )
            valores_arquivo_2[chave] = linha["Valor no Arquivo 2"]

        quantidade_alteracoes = 0

        # --------------------------------------------------------
        # APLICA AS DECISÕES DA INTERFACE
        # --------------------------------------------------------

        for _, decisao in decisoes.iterrows():
            escolha = decisao["Decisão"]

            if escolha != "Usar Arquivo 2":
                continue

            linha_excel = int(decisao["Índice/Linha"])
            nome_coluna = str(decisao["Nome da Coluna"])

            if nome_coluna not in mapa_colunas:
                raise ValueError(
                    f"A coluna '{nome_coluna}' existe no Arquivo 2, "
                    "mas não existe no Arquivo 1. "
                    "Não é possível substituir essa célula."
                )

            chave = (linha_excel, nome_coluna)

            if chave not in valores_arquivo_2:
                raise ValueError(
                    "Não foi possível localizar o valor do Arquivo 2 "
                    f"para a célula {nome_coluna} / linha {linha_excel}."
                )

            novo_valor = normalizar_valor_excel(
                valores_arquivo_2[chave]
            )

            worksheet.cell(
                row=linha_excel,
                column=mapa_colunas[nome_coluna],
            ).value = novo_valor

            quantidade_alteracoes += 1

        # --------------------------------------------------------
        # ADICIONA AUTOMATICAMENTE AS LINHAS NOVAS DO ARQUIVO 2
        # --------------------------------------------------------

        quantidade_linhas_arquivo_1 = len(df_original)
        quantidade_linhas_arquivo_2 = len(df_modificado)

        quantidade_linhas_novas = max(
            0,
            quantidade_linhas_arquivo_2 - quantidade_linhas_arquivo_1,
        )

        if quantidade_linhas_novas > 0:

            # Os dados começam na linha 2, pois a linha 1 é o cabeçalho.
            primeira_linha_nova = quantidade_linhas_arquivo_1 + 2

            # Preserva a formatação da última linha existente sempre que
            # possível.
            ultima_linha_existente = quantidade_linhas_arquivo_1 + 1

            for deslocamento, indice_df2 in enumerate(
                range(quantidade_linhas_arquivo_1, quantidade_linhas_arquivo_2)
            ):

                linha_excel = primeira_linha_nova + deslocamento

                copiar_estilo_linha(
                    worksheet,
                    ultima_linha_existente,
                    linha_excel,
                )

                # Grava somente as colunas que já existem no Arquivo 1,
                # preservando a estrutura do arquivo original.
                for nome_coluna, coluna_excel in mapa_colunas.items():

                    if nome_coluna not in df_modificado.columns:
                        continue

                    valor = df_modificado.iloc[
                        indice_df2
                    ][nome_coluna]

                    worksheet.cell(
                        row=linha_excel,
                        column=coluna_excel,
                    ).value = normalizar_valor_excel(valor)

        saida = BytesIO()
        workbook.save(saida)
        saida.seek(0)

        return (
            saida.getvalue(),
            quantidade_alteracoes,
            quantidade_linhas_novas,
        )

    finally:
        workbook.close()


# ============================================================
# IDENTIFICADOR DA COMPARAÇÃO
# ============================================================

def gerar_id_comparacao(
    arquivo_1_bytes: bytes,
    arquivo_2_bytes: bytes,
    sheet_name: str,
) -> str:
    """Gera um identificador estável para a comparação atual."""

    conteudo = (
        arquivo_1_bytes
        + arquivo_2_bytes
        + sheet_name.encode("utf-8")
    )

    return hashlib.md5(conteudo).hexdigest()


# ============================================================
# INTERFACE
# ============================================================

st.markdown(
    '<div class="main-title">📊 Auditoria de Planilhas Excel</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="subtitle">'
    'Compare duas versões de uma planilha, escolha o valor correto '
    'para cada diferença e gere uma versão corrigida.'
    '</div>',
    unsafe_allow_html=True,
)


# ============================================================
# UPLOAD DOS ARQUIVOS
# ============================================================

col1, col2 = st.columns(2)

with col1:
    arquivo_1 = st.file_uploader(
        "📁 Arquivo 1 — Original",
        type=["xlsx", "xlsm"],
        key="arquivo_1",
    )

with col2:
    arquivo_2 = st.file_uploader(
        "📁 Arquivo 2 — Modificado",
        type=["xlsx", "xlsm"],
        key="arquivo_2",
    )


if arquivo_1 is None or arquivo_2 is None:
    st.info(
        "Selecione os dois arquivos Excel para iniciar a comparação."
    )
    st.stop()


# ============================================================
# CONVERSÃO PARA BYTES
# ============================================================

arquivo_1_bytes = arquivo_1.getvalue()
arquivo_2_bytes = arquivo_2.getvalue()


# ============================================================
# VERIFICAÇÃO DAS ABAS
# ============================================================

try:
    abas_1 = obter_abas(arquivo_1_bytes)
    abas_2 = obter_abas(arquivo_2_bytes)

except Exception as erro:
    st.error(f"Não foi possível abrir os arquivos: {erro}")
    st.stop()


abas_comuns = [
    aba for aba in abas_1 if aba in abas_2
]

if not abas_comuns:
    st.error(
        "Os arquivos não possuem nenhuma aba com o mesmo nome."
    )
    st.warning(
        "Abas do Arquivo 1: " + ", ".join(abas_1) + "\n\n"
        "Abas do Arquivo 2: " + ", ".join(abas_2)
    )
    st.stop()


# ============================================================
# SELEÇÃO DA ABA
# ============================================================

col_aba, col_info = st.columns([2, 4])

with col_aba:
    sheet_name = st.selectbox(
        "📄 Aba para comparar",
        options=abas_comuns,
        key="aba_comparacao",
    )

with col_info:
    abas_somente_1 = [
        aba for aba in abas_1 if aba not in abas_2
    ]
    abas_somente_2 = [
        aba for aba in abas_2 if aba not in abas_1
    ]

    if abas_somente_1 or abas_somente_2:
        mensagem = "⚠️ Diferenças de estrutura entre abas:\n"

        if abas_somente_1:
            mensagem += (
                "- Somente no Arquivo 1: "
                + ", ".join(abas_somente_1)
                + "\n"
            )

        if abas_somente_2:
            mensagem += (
                "- Somente no Arquivo 2: "
                + ", ".join(abas_somente_2)
            )

        st.warning(mensagem)


# ============================================================
# COMPARAÇÃO
# ============================================================

with st.spinner("Comparando as planilhas..."):
    try:
        relatorio, quantidade_linhas_novas = comparar_dataframes(
            arquivo_1_bytes,
            arquivo_2_bytes,
            sheet_name,
        )

    except Exception as erro:
        st.error(f"Erro durante a comparação: {erro}")
        st.stop()


# A comparação pode não ter nenhuma diferença célula a célula e ainda
# possuir novas linhas no Arquivo 2. Nesse caso, a geração continua
# normalmente e as novas linhas serão adicionadas automaticamente.
if relatorio.empty and quantidade_linhas_novas == 0:
    st.success(
        "✅ Nenhuma diferença foi encontrada nesta aba."
    )
    st.stop()

if quantidade_linhas_novas > 0:
    st.info(
        f"➕ {quantidade_linhas_novas} linha(s) existente(s) somente no "
        "Arquivo 2 serão adicionada(s) automaticamente ao arquivo corrigido. "
        "Essas linhas não aparecem na tabela de decisões."
    )


# ============================================================
# CONFIGURAÇÃO DA TABELA
# ============================================================

tabela_inicial = relatorio.copy()
quantidade_total = len(tabela_inicial)

# Chave muda quando os arquivos ou a aba mudam.
# Isso impede que decisões antigas sejam reaproveitadas.
id_comparacao = gerar_id_comparacao(
    arquivo_1_bytes,
    arquivo_2_bytes,
    sheet_name,
)

chave_tabela = f"tabela_diferencas_{id_comparacao}"


# ============================================================
# INSTRUÇÕES / TABELA EDITÁVEL
# ============================================================

if not tabela_inicial.empty:

    st.markdown(
        """
        ### 🔎 Decida cada diferença

        Para cada célula alterada, escolha:

        **Manter Arquivo 1** → mantém o valor original.

        **Usar Arquivo 2** → substitui o valor do Arquivo 1 pelo
        valor existente no Arquivo 2.

        As linhas que existem somente no Arquivo 2 são adicionadas
        automaticamente e não precisam de decisão manual.
        """
    )

    tabela_editada = st.data_editor(
        tabela_inicial,
        key=chave_tabela,
        hide_index=True,
        width="stretch",
        height=650,
        row_height=55,
        disabled=[
            "Índice/Linha",
            "Nome da Coluna",
            "Valor no Arquivo 1",
            "Valor no Arquivo 2",
        ],
        column_config={
            "Índice/Linha": st.column_config.NumberColumn(
                "Linha",
                width="small",
            ),
            "Nome da Coluna": st.column_config.TextColumn(
                "Coluna",
                width="medium",
            ),
            "Valor no Arquivo 1": st.column_config.TextColumn(
                "Arquivo 1 — Original",
                width="large",
            ),
            "Valor no Arquivo 2": st.column_config.TextColumn(
                "Arquivo 2 — Modificado",
                width="large",
            ),
            "Decisão": st.column_config.SelectboxColumn(
                "✅ Decisão",
                help=(
                    "Escolha qual valor deverá permanecer "
                    "no arquivo corrigido."
                ),
                options=[
                    "Pendente",
                    "Manter Arquivo 1",
                    "Usar Arquivo 2",
                ],
                required=True,
                width="medium",
            ),
        },
    )

else:
    tabela_editada = pd.DataFrame(
        columns=[
            "Índice/Linha",
            "Nome da Coluna",
            "Valor no Arquivo 1",
            "Valor no Arquivo 2",
            "Decisão",
        ]
    )

    st.success(
        "✅ Não há diferenças de células para decidir manualmente."
    )


# ============================================================
# CONTADORES
# ============================================================

quantidade_pendente = int(
    tabela_editada["Decisão"].eq("Pendente").sum()
)

usar_arquivo_2_manual = int(
    tabela_editada["Decisão"].eq("Usar Arquivo 2").sum()
)

manter_arquivo_1 = int(
    tabela_editada["Decisão"].eq("Manter Arquivo 1").sum()
)

# Linhas novas não passam pela tabela: são automaticamente consideradas
# como "Usar Arquivo 2".
usar_arquivo_2 = usar_arquivo_2_manual + quantidade_linhas_novas


# ============================================================
# MÉTRICAS
# ============================================================

st.divider()

m1, m2, m3, m4 = st.columns(4)

with m1:
    st.metric("Diferenças", quantidade_total)

with m2:
    st.metric("Pendentes", quantidade_pendente)

with m3:
    st.metric("Usar Arquivo 2", usar_arquivo_2)

with m4:
    st.metric("Manter Arquivo 1", manter_arquivo_1)


# ============================================================
# RESUMO E BOTÃO
# ============================================================

col_resumo, col_botao = st.columns([3, 2])

with col_resumo:
    st.markdown(
        f"""
        **Resumo das decisões**

        Total de diferenças: **{quantidade_total}**

        Pendentes: **{quantidade_pendente}**

        Manter Arquivo 1: **{manter_arquivo_1}**

        Usar Arquivo 2: **{usar_arquivo_2}**

        Inclusões automáticas: **{quantidade_linhas_novas} linha(s)**
        """
    )

with col_botao:
    gerar = st.button(
        "🚀 GERAR ARQUIVO CORRIGIDO",
        type="primary",
        use_container_width=True,
        key=f"gerar_{id_comparacao}",
    )


# ============================================================
# GERAÇÃO DO ARQUIVO
# ============================================================

if gerar:
    if quantidade_pendente > 0:
        st.error(
            f"Existem {quantidade_pendente} diferença(s) sem decisão."
        )
        st.warning(
            "Escolha 'Manter Arquivo 1' ou 'Usar Arquivo 2' "
            "para todas as diferenças."
        )
        st.stop()

    with st.spinner("Gerando arquivo corrigido..."):
        try:
            (
                arquivo_corrigido,
                quantidade_alteracoes,
                quantidade_linhas_adicionadas,
            ) = gerar_arquivo_corrigido(
                arquivo_1_bytes=arquivo_1_bytes,
                arquivo_2_bytes=arquivo_2_bytes,
                relatorio_original=relatorio,
                decisoes=tabela_editada,
                sheet_name=sheet_name,
                nome_arquivo_original=arquivo_1.name,
            )

        except Exception as erro:
            st.error(
                f"Erro ao gerar o arquivo corrigido: {erro}"
            )
            st.stop()

    nome_original = Path(arquivo_1.name)

    nome_saida = (
        nome_original.stem
        + "_corrigido"
        + nome_original.suffix
    )

    mime = (
        "application/vnd.ms-excel.sheet.macroEnabled.12"
        if nome_original.suffix.lower() == ".xlsm"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    st.success(
        "✅ Arquivo corrigido gerado com sucesso! "
        f"{quantidade_alteracoes} célula(s) foram alteradas e "
        f"{quantidade_linhas_adicionadas} linha(s) nova(s) foram adicionadas automaticamente."
    )

    st.download_button(
        label="⬇️ BAIXAR ARQUIVO CORRIGIDO",
        data=arquivo_corrigido,
        file_name=nome_saida,
        mime=mime,
        type="primary",
        use_container_width=True,
        key=f"download_{id_comparacao}_{quantidade_alteracoes}",
    )
